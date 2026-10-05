# -*- coding: utf-8 -*-
"""Формирует Excel-отчёт по результатам аудита: AUDIT_OTCHET.xlsx"""
import json, sys, collections
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

res = json.load(open('audit_results.json', encoding='utf-8'))
code2name = json.load(open('code2name.json', encoding='utf-8'))

RJU = {1: 'Ташкентский', 2: 'Кокандский (Ферганское)', 3: 'Бухарский', 4: 'Кунградский', 5: 'Каршинский', 6: 'Термезский'}
STRANA = {'20': 'РЖД', '21': 'БЧ (Беларусь)', '22': 'УЗ (Украина)', '23': 'ЧФМ (Молдова)', '24': 'ЛГ (Литва)',
          '25': 'ЛДЗ (Латвия)', '26': 'ЭВР (Эстония)', '27': 'КТЖ (Казахстан)', '28': 'ГР (Грузия)', '29': 'УТИ (Узбекистан)',
          '57': 'АДЙ (Азербайджан)', '58': 'ЮКЖД (Армения)', '59': 'КРГ (Кыргызстан)', '66': 'ТДЖ (Таджикистан)', '67': 'ТРК (Туркменистан)'}

def chain_str(codes):
    return ' → '.join(f"{code2name.get(c, '?')} ({c})" for c in codes)

wb = Workbook()
hdr_font = Font(bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='305496')
red = PatternFill('solid', fgColor='FFC7CE')
yellow = PatternFill('solid', fgColor='FFEB9C')
green = PatternFill('solid', fgColor='C6EFCE')

def sheet(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, i, h); c.font = hdr_font; c.fill = hdr_fill
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

def rowbase(r):
    ro = r['row']
    return [r['id'], RJU.get(ro['sendRjuId'], ro['sendRjuId']),
            f"{ro['sendStationBegin']}-{ro['sendStationEnd']}",
            STRANA.get(ro['receiveStranaCode'], ro['receiveStranaCode']),
            f"{ro['receiveStationBegin']}-{ro['receiveStationEnd']}"]

# ---- Xulosa ----
ws = wb.active; ws.title = 'Сводка'
cnt = collections.Counter(r['status'] for r in res)
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 10; ws.column_dimensions['C'].width = 95
ws.cell(1, 1, 'Аудит маршрутных планов e-nakl по Книге 3 (2025-10-23)').font = Font(bold=True, size=14)
ws.cell(2, 1, f"Всего проверено: {len(res)} записей")
rows_meta = [
    ('OK', 'Полностью соответствует Книге 3 — можно подтверждать', green),
    ('OK_RANGE_MISALIGNED', 'Пункты перехода верны, но диапазон станций не совпадает с разбивкой книги', green),
    ('CHAIN_BY_NOTE_PARTIAL', 'Примечание книги допускает этот маршрут, но на сайте пропущены промежуточные пункты', yellow),
    ('CHAIN_MISMATCH', 'ОШИБКА: цепочка пунктов перехода не соответствует книге', red),
    ('DEST_RANGE_AMBIGUOUS', 'Диапазон пересекает 2+ строки книги с разными цепочками — запись нужно разделить', red),
    ('DEST_NOT_FOUND', 'Диапазон приёма вообще не найден в книге', red),
    ('BAD_INPUT_DATA', 'Грубая ошибка ввода (обратный диапазон, 6-значный код, дубликат пункта)', red),
    ('DOMESTIC_NO_RULE_NEEDED', 'Чисто внутренний маршрут — правило Книги 3 не требуется', green),
]
r0 = 4
ws.cell(r0, 1, 'Категория').font = Font(bold=True)
ws.cell(r0, 2, 'Кол-во').font = Font(bold=True)
ws.cell(r0, 3, 'Пояснение').font = Font(bold=True)
for i, (st, desc, fill) in enumerate(rows_meta, 1):
    ws.cell(r0 + i, 1, st).fill = fill
    ws.cell(r0 + i, 2, cnt.get(st, 0))
    ws.cell(r0 + i, 3, desc)

# ---- Zanjir xatolari ----
ws = wb.create_sheet('Ошибки цепочек')
sheet(ws, ['ID', 'Узел', 'Отправление', 'Дорога', 'Приём', 'Цепочка на сайте', 'Цепочка по книге', 'Расхождение (позиция)'],
      [7, 16, 14, 12, 14, 55, 55, 40])
i = 2
for r in res:
    if r['status'] not in ('CHAIN_MISMATCH', 'CHAIN_BY_NOTE_PARTIAL'): continue
    d = r['detail']
    if r['status'] == 'CHAIN_BY_NOTE_PARTIAL':
        diffs = "Допущено примечанием; полная цепочка из примечания: " + chain_str(d.get('note_chain') or [])
    else:
        diffs = '; '.join(f"{x['pos']}: {code2name.get(x['actual'] or '', '—')}({x['actual'] or '—'}) ≠ {code2name.get(x['expected'] or '', '—')}({x['expected'] or '—'})"
                          for x in d.get('diffs', []))
    for j, v in enumerate(rowbase(r) + [chain_str(d['actual_chain']), chain_str(d.get('expected_chain') or []), diffs], 1):
        ws.cell(i, j, v)
    ws.cell(i, 1).fill = yellow if r['status'] == 'CHAIN_BY_NOTE_PARTIAL' else red
    i += 1

# ---- Diapazon muammolari ----
ws = wb.create_sheet('Проблемы диапазонов')
sheet(ws, ['ID', 'Узел', 'Отправление', 'Дорога', 'Приём', 'Статус', 'Цепочка на сайте', 'Пояснение'],
      [7, 16, 14, 12, 14, 24, 50, 70])
i = 2
for r in res:
    if r['status'] not in ('DEST_RANGE_AMBIGUOUS', 'DEST_NOT_FOUND', 'OK_RANGE_MISALIGNED'): continue
    d = r['detail']
    note = ''
    if r['status'] == 'DEST_RANGE_AMBIGUOUS':
        note = 'Пересекаемые строки книги: ' + ' | '.join(
            f"{road[:35]}: {'-'.join(ch)}" for _, road, ch in d.get('intersecting_rules', [])[:3])
    elif r['status'] == 'OK_RANGE_MISALIGNED':
        note = f"Цепочка верна ({d.get('dest_road','')[:40]}), приведите диапазон к разбивке книги"
    for j, v in enumerate(rowbase(r) + [r['status'], chain_str(d['actual_chain']), note], 1):
        ws.cell(i, j, v)
    ws.cell(i, 1).fill = yellow if r['status'] == 'OK_RANGE_MISALIGNED' else red
    i += 1

# ---- Qo'pol xatolar ----
ws = wb.create_sheet('Грубые ошибки')
sheet(ws, ['ID', 'Узел', 'Отправление', 'Дорога', 'Приём', 'Пункты (как введено на сайте)', 'Проблема'],
      [7, 16, 16, 12, 16, 40, 60])
i = 2
for r in res:
    if r['status'] != 'BAD_INPUT_DATA': continue
    for j, v in enumerate(rowbase(r) + [r['row']['punktsConcat'], ', '.join(r['detail']['flags'])], 1):
        ws.cell(i, j, v)
    ws.cell(i, 1).fill = red
    i += 1

# ---- OK ro'yxati (bulk tasdiqlash uchun) ----
ws = wb.create_sheet('Список OK')
sheet(ws, ['ID', 'Узел', 'Отправление', 'Дорога', 'Приём', 'Цепочка', 'Статус'], [7, 18, 14, 18, 14, 55, 20])
i = 2
for r in res:
    if r['status'] not in ('OK',): continue
    for j, v in enumerate(rowbase(r) + [chain_str(r['detail']['actual_chain']), 'OK'], 1):
        ws.cell(i, j, v)
    i += 1

wb.save('AUDIT_OTCHET.xlsx')
print('saved AUDIT_OTCHET.xlsx')
